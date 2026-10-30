"""Read an authenticated Drive XLSX download; validate and emit an atomic SQL import.
Usage: python parse-jv-workbook.py source.xlsx cutoff_iso modified_iso > import.sql
Never executes SQL or changes the source workbook.
"""
import sys,json,re,datetime
from decimal import Decimal,ROUND_HALF_UP
import openpyxl

def decimal(value):
    if isinstance(value,bool) or not isinstance(value,(int,float,Decimal)):
        raise ValueError('Montant absent, texte ou erreur Excel : '+repr(value))
    d=Decimal(str(value))
    if not d.is_finite(): raise ValueError('Montant non fini')
    return d
def rounded(d): return d.quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
def equal(a,b):
    if abs(decimal(a)-decimal(b))>Decimal('.005'):raise ValueError(f'Calcul incohérent : {a} au lieu de {b}')
def parse(path,modified):
    wb=openpyxl.load_workbook(path,data_only=True)
    tariffs=wb['Tarifs garanties']
    costs={tariffs.cell(r,1).value:decimal(tariffs.cell(r,3).value) for r in range(2,9)}
    names=['Janvier','Février','Mars','Avril','Mai','Juin','Juillet','Août','Septembre','Octobre','Novembre','Décembre']
    result=[];seen=set()
    for s in wb:
        match=re.fullmatch(r'('+ '|'.join(names)+r') (20\d{2})',s.title)
        if not match:continue
        month=names.index(match[1])+1;year=int(match[2]);seen.add((year,month))
        rate=decimal(s['B5'].value);vat=decimal(s['D5'].value);cash=decimal(s['F5'].value);share=decimal(s['H5'].value)
        if rate!=Decimal('.08') or not 0<=vat<=1 or not 0<=share<=1:raise ValueError('Paramètres à vérifier')
        totals_row=next((r for r in range(8,s.max_row+1) if s.cell(r,7).value=='TOTAL À FACTURER HT'),None)
        if not totals_row:raise ValueError('Total introuvable : '+s.title)
        rows=[]
        for r in range(8,totals_row):
            # Notes after the detail range are not sales.
            if s.cell(r,1).value and str(s.cell(r,1).value).startswith('Base HT ='):continue
            inputs=[s.cell(r,c).value for c in (1,2,3,5)]
            if all(v in (None,'') for v in inputs):continue
            commercial,vehicle,margin,guarantee=inputs
            if margin in (None,'') or not str(commercial or '').strip() or not str(vehicle or '').strip():raise ValueError(f'{s.title} ligne {r} : saisie incomplète')
            margin=decimal(margin);mh=rounded(margin/(1+vat));gh=costs[guarantee] if guarantee else Decimal(0)
            ch=rounded(cash/(1+vat)*share);base=rounded(mh-gh-ch);commission=rounded(base*rate)
            for col,expected in ((4,mh),(6,ch),(7,base),(8,commission)):equal(s.cell(r,col).value,expected)
            rows.append(dict(row=r,commercial=str(commercial).strip(),vehicle=str(vehicle).strip(),margin_ttc=float(margin),margin_ht=float(mh),guarantee=guarantee or '',guarantee_ht=float(gh),cash_ht=float(ch),base_ht=float(base),commission_ht=float(commission)))
        total=sum((Decimal(str(r['commission_ht'])) for r in rows),Decimal(0));tax=rounded(total*vat)
        for offset,expected in ((0,total),(1,tax),(2,total+tax)):equal(s.cell(totals_row+offset,8).value,expected)
        result.append(dict(month_start=f'{year}-{month:02d}-01',source_modified_at=modified,total_ht=float(total),total_vat=float(tax),total_ttc=float(total+tax),rows=rows))
    if not result:raise ValueError('Aucun onglet mensuel reconnu')
    for year in {y for y,m in seen}:
        if {(year,m) for m in range(1,13)}-seen:raise ValueError('Année incomplète : onglets manquants')
    return result

if __name__=='__main__':
    data=parse(sys.argv[1],sys.argv[3]);cutoff=datetime.datetime.fromisoformat(sys.argv[2].replace('Z','+00:00'))
    if cutoff.tzinfo is None:raise ValueError('Fuseau horaire absent')
    payload=json.dumps(data,ensure_ascii=False,allow_nan=False).replace("'","''")
    print("select private.import_tours_jv('"+payload+"'::jsonb,'"+cutoff.isoformat()+"'::timestamptz);")
